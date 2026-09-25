"""T-020 vault/fill tests using synthetic values and an in-memory keyring."""
from __future__ import annotations

import contextlib
import io
import json
import tempfile
import unittest
import unittest.mock
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from pyrage import x25519

from pipeline import config, db, fill, settings, vault


class MemoryCredentialStore:
    def __init__(self) -> None:
        self.values: dict[tuple[str, str], str] = {}

    def get_password(self, service: str, username: str) -> str | None:
        return self.values.get((service, username))

    def set_password(self, service: str, username: str, password: str) -> None:
        self.values[(service, username)] = password

    def delete_password(self, service: str, username: str) -> None:
        self.values.pop((service, username), None)


class FakeLocator:
    def __init__(self, page: "FakePage", selector: str) -> None:
        self.page = page
        self.selector = selector
        self.masked = False
        self._value = ""

    def evaluate(self, _script: str) -> None:
        self.masked = True

    def fill(self, value: str) -> None:
        self._value = value


class FakePage:
    def __init__(self, url: str = "https://jobs.example.test/apply") -> None:
        self.url = url
        self.locators: dict[str, FakeLocator] = {}
        self.styles: list[str] = []

    def add_style_tag(self, *, content: str) -> None:
        self.styles.append(content)

    def locator(self, selector: str) -> FakeLocator:
        self.locators.setdefault(selector, FakeLocator(self, selector))
        return self.locators[selector]

    def screenshot(self, *, path: str, full_page: bool) -> None:
        assert full_page
        rendered = "\n".join(
            f"{selector}={'[MASKED]' if locator.masked else locator._value}"
            for selector, locator in sorted(self.locators.items())
        )
        Path(path).write_text(rendered, encoding="utf-8")


class VaultTestCase(unittest.TestCase):
    SECRET = "SYNTHETIC_SECRET_NEVER_DISCLOSE_7f9a"

    def setUp(self) -> None:
        self.temp_context = tempfile.TemporaryDirectory(prefix="careeros-vault-test-")
        self.root = Path(self.temp_context.name)
        self.old_paths = {
            name: getattr(config, name)
            for name in (
                "LOCAL_DIR", "DB_PATH", "BACKUP_DIR", "LOGS_DIR", "SECRET_VAULT_DIR",
                "INBOX_DIR", "SCREENSHOTS_DIR", "PAUSED_FLAG",
            )
        }
        config.LOCAL_DIR = self.root
        config.DB_PATH = self.root / "tracker.sqlite"
        config.BACKUP_DIR = self.root / "backups"
        config.LOGS_DIR = self.root / "logs"
        config.SECRET_VAULT_DIR = self.root / "vault"
        config.INBOX_DIR = self.root / "inbox"
        config.SCREENSHOTS_DIR = self.root / "screenshots"
        config.PAUSED_FLAG = self.root / "PAUSED"
        config.ensure_dirs()
        self.conn = db.connect(config.DB_PATH)
        db.migrate(self.conn)
        self.credentials = MemoryCredentialStore()
        self.vault = vault.Vault(
            config.SECRET_VAULT_DIR,
            credential_store=self.credentials,
            enforce_windows_backend=False,
            enforce_acl=False,
            now=lambda: "2026-09-22T12:00:00+00:00",
        )
        self.auto_settings = settings.Settings(
            data={
                "autonomy": {"overrides": {"fill_identity_fields": "auto"}},
                "vault": {
                    "password_policy": "unique",
                    "applications_email": "applications@example.test",
                    "autofill_domains_allowlist": ["example.test"],
                },
            }
        )

    def tearDown(self) -> None:
        self.conn.close()
        for name, value in self.old_paths.items():
            setattr(config, name, value)
        self.temp_context.cleanup()

    def test_age_round_trip_index_and_ciphertext_never_contain_value(self) -> None:
        self.vault.set(
            "identity.phone",
            self.SECRET,
            metadata={"kind": "identity", "field": "phone"},
        )
        ciphertext = next(config.SECRET_VAULT_DIR.glob("*.age")).read_bytes()
        index = (config.SECRET_VAULT_DIR / "index.json").read_text(encoding="utf-8")
        self.assertNotIn(self.SECRET.encode(), ciphertext)
        self.assertNotIn(self.SECRET, index)
        self.assertTrue(self.vault.contains("identity.phone"))
        self.assertEqual(self.vault._read_value("identity.phone"), self.SECRET)

    def test_ciphertext_is_random_and_bound_to_its_record_key(self) -> None:
        self.vault.set("identity.one", self.SECRET)
        first = self.vault._record_path("identity.one").read_bytes()
        self.vault.set("identity.one", self.SECRET)
        second = self.vault._record_path("identity.one").read_bytes()
        self.assertNotEqual(first, second)
        self.vault.set("identity.two", "different synthetic value")
        self.vault._record_path("identity.two").write_bytes(second)
        with self.assertRaisesRegex(vault.VaultError, "invalid"):
            self.vault._read_value("identity.two")

    def test_decrypted_metadata_must_match_index(self) -> None:
        self.vault.set(
            "identity.phone",
            self.SECRET,
            metadata={"kind": "identity", "field": "phone"},
        )
        index_path = config.SECRET_VAULT_DIR / "index.json"
        index = json.loads(index_path.read_text(encoding="utf-8"))
        index["records"]["identity.phone"]["field"] = "address"
        index_path.write_text(json.dumps(index), encoding="utf-8")
        with self.assertRaisesRegex(vault.VaultError, "invalid"):
            self.vault._read_value("identity.phone")

    def test_wrong_identity_fails_with_redacted_error(self) -> None:
        self.vault.set("identity.phone", self.SECRET)
        other = x25519.Identity.generate()
        self.credentials.values[(vault.KEYRING_SERVICE, vault.KEYRING_ACCOUNT)] = str(other)
        try:
            self.vault._read_value("identity.phone")
        except vault.VaultError as exc:
            message = str(exc)
        else:
            self.fail("wrong identity unexpectedly decrypted a record")
        self.assertNotIn(self.SECRET, message)
        self.assertNotIn(str(other), message)

    def test_portal_password_policies_never_return_passwords(self) -> None:
        unique = vault.prepare_portal_credentials(
            "https://jobs.example.test/path", vault=self.vault, active_settings=self.auto_settings
        )
        self.assertEqual(unique.site, "jobs.example.test")
        self.assertNotIn(self.SECRET, repr(unique))
        password = self.vault._read_value(unique.password_key)
        self.assertGreaterEqual(len(password), 20)
        self.assertNotEqual(password, "applications@example.test")

        single_settings = settings.Settings(
            data={
                "vault": {
                    "password_policy": "single",
                    "applications_email": "applications@example.test",
                }
            }
        )
        first = vault.prepare_portal_credentials(
            "first.example", vault=self.vault, active_settings=single_settings
        )
        second = vault.prepare_portal_credentials(
            "second.example", vault=self.vault, active_settings=single_settings
        )
        self.assertTrue(
            self.vault._read_value(first.password_key)
            == self.vault._read_value(second.password_key)
        )

    def test_fill_masks_before_value_and_audit_contains_only_key_metadata(self) -> None:
        self.vault.set("identity.phone", self.SECRET)
        page = FakePage()
        screenshot = config.SCREENSHOTS_DIR / "masked.png"
        output = io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(output):
            result = fill.fill_fields(
                page,
                "https://jobs.example.test/apply?token=not-logged",
                [fill.FieldFromVault("#phone", "phone", "identity.phone")],
                vault=self.vault,
                active_settings=self.auto_settings,
                screenshot_path=screenshot,
                conn=self.conn,
            )
        self.assertEqual(result.site, "jobs.example.test")
        self.assertTrue(page.locators["#phone"].masked)
        self.assertNotIn(self.SECRET, screenshot.read_text(encoding="utf-8"))
        event = self.conn.execute(
            "SELECT detail_json FROM events WHERE type='field_filled' ORDER BY id DESC LIMIT 1"
        ).fetchone()[0]
        self.assertEqual(
            json.loads(event),
            {"site": "jobs.example.test", "field": "phone", "key": "identity.phone"},
        )
        combined = output.getvalue() + event + screenshot.read_text(encoding="utf-8") + repr(result)
        self.assertNotIn(self.SECRET, combined)
        self.assertNotIn("not-logged", combined)

    def test_real_chromium_fill_is_password_typed_and_visually_masked(self) -> None:
        self.vault.set("identity.phone", self.SECRET)
        screenshot = config.SCREENSHOTS_DIR / "chromium-masked.png"
        output = io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(output):
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                try:
                    context = browser.new_context()
                    context.route(
                        "https://jobs.example.test/**",
                        lambda route: route.fulfill(
                            body='<label>Phone <input id="phone" type="tel"></label>',
                            content_type="text/html",
                        ),
                    )
                    page = context.new_page()
                    page.goto("https://jobs.example.test/apply")
                    fill.fill_fields(
                        page,
                        "jobs.example.test",
                        [fill.FieldFromVault("#phone", "phone", "identity.phone")],
                        vault=self.vault,
                        active_settings=self.auto_settings,
                        screenshot_path=screenshot,
                        conn=self.conn,
                    )
                    state = page.locator("#phone").evaluate(
                        "el => ({type: el.type, masked: el.dataset.careerosSecret, "
                        "textSecurity: getComputedStyle(el).webkitTextSecurity})"
                    )
                    # What Playwright MCP would hand the model, before and after
                    # the field types are restored for validation.
                    snapshots = [page.locator("body").aria_snapshot()]
                    fill.restore_field_types_for_validation(page)
                    snapshots.append(page.locator("body").aria_snapshot())
                finally:
                    browser.close()
        self.assertEqual(state["type"], "password")
        self.assertEqual(state["masked"], "true")
        self.assertEqual(state["textSecurity"], "disc")
        self.assertTrue(screenshot.is_file())
        self.assertNotIn(self.SECRET, output.getvalue())
        for snapshot in snapshots:
            self.assertNotIn(self.SECRET, snapshot)

    def test_portal_credential_only_fills_on_its_own_site(self) -> None:
        keys = vault.prepare_portal_credentials(
            "a.example.test", vault=self.vault, active_settings=self.auto_settings
        )
        page = FakePage("https://b.example.test/login")
        with self.assertRaisesRegex(fill.FillError, "bound to another site"):
            fill.fill_fields(
                page,
                "b.example.test",
                [fill.FieldFromVault("#password", "password", keys.password_key)],
                vault=self.vault,
                active_settings=self.auto_settings,
                conn=self.conn,
            )
        self.assertEqual(page.locators["#password"]._value, "")
        result = fill.fill_fields(
            FakePage("https://a.example.test/login"),
            "a.example.test",
            [fill.FieldFromVault("#password", "password", keys.password_key)],
            vault=self.vault,
            active_settings=self.auto_settings,
            conn=self.conn,
        )
        self.assertEqual(result.fields, ("password",))

    def test_redirect_between_fields_stops_the_next_fill(self) -> None:
        self.vault.set("identity.phone", self.SECRET)
        self.vault.set("identity.address", self.SECRET + "-address")
        page = FakePage()
        original_fill = FakeLocator.fill

        def fill_then_redirect(locator: FakeLocator, value: str) -> None:
            original_fill(locator, value)
            page.url = "https://attacker.example/phish"

        FakeLocator.fill = fill_then_redirect  # type: ignore[method-assign]
        try:
            with self.assertRaisesRegex(fill.FillError, "does not match"):
                fill.fill_fields(
                    page,
                    "jobs.example.test",
                    [
                        fill.FieldFromVault("#phone", "phone", "identity.phone"),
                        fill.FieldFromVault("#address", "address", "identity.address"),
                    ],
                    vault=self.vault,
                    active_settings=self.auto_settings,
                    conn=self.conn,
                )
        finally:
            FakeLocator.fill = original_fill  # type: ignore[method-assign]
        self.assertNotIn("#address", page.locators)

    def test_pause_and_playwright_debug_block_automatic_fills(self) -> None:
        self.vault.set("identity.phone", self.SECRET)
        mapping = [fill.FieldFromVault("#phone", "phone", "identity.phone")]
        config.PAUSED_FLAG.write_text("paused", encoding="utf-8")
        try:
            with self.assertRaises(fill.ApprovalRequired):
                fill.fill_fields(
                    FakePage(), "jobs.example.test", mapping,
                    vault=self.vault, active_settings=self.auto_settings, conn=self.conn,
                )
        finally:
            config.PAUSED_FLAG.unlink()
        with unittest.mock.patch.dict("os.environ", {"DEBUG": "pw:api"}):
            with self.assertRaisesRegex(fill.FillError, "debug"):
                fill.fill_fields(
                    FakePage(), "jobs.example.test", mapping,
                    vault=self.vault, active_settings=self.auto_settings,
                    human_confirmed=True, conn=self.conn,
                )

    def test_fill_rejects_browser_page_on_another_domain(self) -> None:
        self.vault.set("identity.phone", self.SECRET)
        page = FakePage("https://lookalike.example.test/apply")
        with self.assertRaisesRegex(fill.FillError, "does not match"):
            fill.fill_fields(
                page,
                "jobs.example.test",
                [fill.FieldFromVault("#phone", "phone", "identity.phone")],
                vault=self.vault,
                active_settings=self.auto_settings,
                conn=self.conn,
            )
        self.assertEqual(page.locators, {})

    def test_fill_requires_allowlist_or_human_confirmation(self) -> None:
        self.vault.set("identity.phone", self.SECRET)
        with self.assertRaises(fill.ApprovalRequired):
            fill.fill_fields(
                FakePage("https://other.example/apply"),
                "other.example",
                [fill.FieldFromVault("#phone", "phone", "identity.phone")],
                vault=self.vault,
                active_settings=self.auto_settings,
                conn=self.conn,
            )
        result = fill.fill_fields(
            FakePage("https://other.example/apply"),
            "other.example",
            [fill.FieldFromVault("#phone", "phone", "identity.phone")],
            vault=self.vault,
            active_settings=self.auto_settings,
            human_confirmed=True,
            conn=self.conn,
        )
        self.assertEqual(result.site, "other.example")

    def test_portal_account_form_respects_create_accounts_override(self) -> None:
        page = FakePage()
        with self.assertRaises(fill.ApprovalRequired):
            fill.fill_portal_account_form(
                page,
                "jobs.example.test",
                email_selector="#email",
                password_selector="#password",
                vault=self.vault,
                active_settings=self.auto_settings,
                conn=self.conn,
            )
        result = fill.fill_portal_account_form(
            page,
            "jobs.example.test",
            email_selector="#email",
            password_selector="#password",
            vault=self.vault,
            active_settings=self.auto_settings,
            human_confirmed=True,
            conn=self.conn,
        )
        self.assertEqual(result.fields, ("account_email", "account_password"))

    def test_verification_code_moves_from_local_email_body_to_masked_field(self) -> None:
        body_path = config.INBOX_DIR / "verification.txt"
        body_path.write_text(
            "Your verification code is 481726 for jobs.example.test.", encoding="utf-8"
        )
        self.conn.execute(
            """
            INSERT INTO emails (
              provider, message_id, received_at, from_addr, from_domain, subject,
              body_path, category, link_domains_json, deadline_confidence
            ) VALUES ('gmail','verify-1','2026-09-22T12:00:00+00:00',
                      'noreply@jobs.example.test','jobs.example.test','Verification code',
                      ?, 'other', '[]', 'unknown')
            """,
            (str(body_path),),
        )
        self.conn.commit()
        page = FakePage()
        screenshot = config.SCREENSHOTS_DIR / "verification.png"
        result = fill.fill_latest_verification_code(
            page,
            "jobs.example.test",
            selector="#code",
            email_conn=self.conn,
            received_after=datetime(2026, 9, 22, 11, 0, tzinfo=timezone.utc),
            active_settings=self.auto_settings,
            screenshot_path=screenshot,
            audit_conn=self.conn,
        )
        self.assertEqual(result.fields, ("verification_code",))
        self.assertTrue(page.locators["#code"].masked)
        self.assertNotIn("481726", screenshot.read_text(encoding="utf-8"))
        event = self.conn.execute(
            "SELECT detail_json FROM events WHERE type='verification_code_filled'"
        ).fetchone()[0]
        self.assertNotIn("481726", event)

    def test_plaintext_export_is_explicit_local_and_audited_without_values(self) -> None:
        keys = vault.prepare_portal_credentials(
            "jobs.example.test", vault=self.vault, active_settings=self.auto_settings
        )
        password = self.vault._read_value(keys.password_key)
        destination = self.vault.export_credentials(conn=self.conn)
        self.assertEqual(destination.parent, config.SECRET_VAULT_DIR)
        workbook = load_workbook(destination, read_only=True)
        try:
            rows = list(workbook["Portal credentials"].iter_rows(values_only=True))
        finally:
            workbook.close()
        self.assertEqual(rows[1][0], "jobs.example.test")
        self.assertTrue(rows[1][3] == password)
        event = self.conn.execute(
            "SELECT detail_json FROM events WHERE type='credentials_exported'"
        ).fetchone()[0]
        self.assertNotIn(password, event)
        self.assertNotIn("applications@example.test", event)


if __name__ == "__main__":
    unittest.main()
