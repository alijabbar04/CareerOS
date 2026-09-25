"""T-002: ingest the candidate's source documents into brain/sources/ as Markdown with front matter.

Inputs (read-only): the <careers-folder> folder (CVs, cover letters, application answers,
speculative emails, tracker, prompts, grades, worksheets, transcript) and the the current employer
work-hours tracker workbook. Originals are never copied; only extracted text is written.

Each output file:  brain/sources/<id>.md   (front matter + text)
Plus:              brain/sources/INDEX.md  and  brain/sources/manifest.json

Re-runnable: unchanged sources are skipped unless --force is given.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

import yaml

from pipeline import config

# ---------------------------------------------------------------------------
# Classification rules
# ---------------------------------------------------------------------------

# Folder (relative, lower-case, forward slashes) -> (kind, voice_corpus)
FOLDER_KINDS = {
    "cvs": ("cv", False),
    "cover letters": ("cover-letter", True),
    "cover letters/(word)": ("cover-letter", True),
    "application questions": ("application-answers", True),
    "application tracker": ("tracker", False),
    "emails": ("speculative-email", True),
    "prompts": ("prompt", False),
    "other": ("other", False),
    "": ("cover-letter", True),   # root-level cover letter
}

SKIP_DIRS = {"job-apply-assistant", "__pycache__", "CareerOS Plan", "CareerOS Backups"}
SKIP_SUFFIXES = {".pyc", ".py", ".yaml", ".yml"}

# Files that are not the candidate's own writing (kept, but excluded from the voice corpus)
NOT_ALI_VOICE = {
    "application questions/prerecorded interview questions.docx": "generic interview-coaching template, not the candidate's writing",
}

EMPLOYER_PATTERNS = [
    re.compile(r"^candidate name c[lv] - (?P<e>.+?)(?: \(nfp\)| all in tax| r&d tax| 2\.0| 20\d\d)?$", re.I),
    re.compile(r"^application questions - (?P<e>.+)$", re.I),
    re.compile(r"^(?P<e>.+?) (?:video interview questions|audit application questions|interview questions)(?: \(\d\))?$", re.I),
    re.compile(r"^why do you want to work at (?P<e>.+)$", re.I),
    re.compile(r"^(?P<e>.+?) application \(aj\)$", re.I),
]


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def docx_text(path: Path) -> tuple[str, dict]:
    from docx import Document

    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    meta = {}
    cp = d.core_properties
    if cp.created:
        meta["doc_created"] = cp.created.date().isoformat()
    if cp.modified:
        meta["doc_modified"] = cp.modified.date().isoformat()
    return "\n".join(parts), meta


def pdf_text(path: Path) -> tuple[str, dict]:
    import fitz  # PyMuPDF

    doc = fitz.open(str(path))
    pages = [page.get_text() for page in doc]
    text = "\n".join(pages)
    meta = {"pages": len(doc)}
    if len(text.strip()) < 40 * max(1, len(doc)):
        meta["needs_ocr"] = True
    return text, meta


def xlsx_text(path: Path) -> tuple[str, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                out.append(" | ".join("" if c is None else str(c).strip() for c in row))
    return "\n".join(out), {}


def plain_text(path: Path) -> tuple[str, dict]:
    return path.read_text(encoding="utf-8", errors="replace"), {}


EXTRACTORS = {".docx": docx_text, ".pdf": pdf_text, ".xlsx": xlsx_text, ".txt": plain_text, ".md": plain_text}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def sha256_text(text: str) -> str:
    return hashlib.sha256(normalise(text).encode("utf-8")).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def slugify(s: str) -> str:
    s = s.lower().replace("&", "and")
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "untitled"


def guess_employer(stem: str) -> str | None:
    for pat in EMPLOYER_PATTERNS:
        m = pat.match(stem.strip())
        if m:
            return m.group("e").strip()
    return None


def kind_for(rel: Path) -> tuple[str, bool, str | None]:
    folder = rel.parent.as_posix().lower()
    if folder == ".":
        folder = ""
    kind, voice = FOLDER_KINDS.get(folder, ("other", False))
    note = NOT_ALI_VOICE.get(rel.as_posix().lower())
    if note:
        voice = False
    name = rel.name.lower()
    if kind == "other":
        if "transcript" in name:
            kind = "transcript"
        elif "grades" in name:
            kind = "grades"
        elif "worksheet" in name:
            kind = "assessment-exercise"
    return kind, voice, note


def front_matter(meta: dict) -> str:
    return "---\n" + yaml.safe_dump(meta, sort_keys=False, allow_unicode=True) + "---\n"


# ---------------------------------------------------------------------------
# Work log (the current employer) — a special-cased source with a day-by-day table
# ---------------------------------------------------------------------------

def ingest_work_log(path: Path) -> dict | None:
    if not path.exists():
        print(f"  ! work log not found: {path}")
        return None
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    rows = []
    for ws in wb.worksheets:
        if ws.title.lower().startswith("summary"):
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            date, day, start, finish, brk, hours, ot, typ, total, pay, summary = (list(row) + [None] * 11)[:11]
            if not isinstance(date, (dt.date, dt.datetime)):
                continue
            summary = (summary or "").replace("•", "-").replace("�", "-").strip()
            if not summary or "no tasks recorded" in summary.lower():
                summary = summary or "(no tasks recorded)"
            rows.append((date.date() if isinstance(date, dt.datetime) else date, typ or "", ot or 0, summary))
    rows.sort()
    lines = ["| Date | Type | Overtime (h) | Tasks recorded |", "|---|---|---|---|"]
    for date, typ, ot, summary in rows:
        lines.append(f"| {date.isoformat()} | {typ} | {ot} | {summary.replace(chr(10), ' ')} |")
    text = "\n".join(lines)
    meta = {
        "id": "work-log--lifted-2026",
        "kind": "work-log",
        "title": "the current employer work-hours tracker, day-by-day task log",
        "employer": "the current employer",
        "role": "Operations Executive",
        "date_range": f"{rows[0][0].isoformat()} to {rows[-1][0].isoformat()}" if rows else None,
        "source_path": str(path),
        "sha256": sha256_file(path),
        "voice_corpus": False,
        "tags": ["lifted", "work-log", "automation", "evidence"],
        "notes": "the candidate's own daily log; the 'Tasks recorded' column is the evidence for the the current employer claims. Days marked '(no tasks recorded)' had no summary entered.",
    }
    return {"meta": meta, "text": text, "words": len(text.split())}


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def ingest(force: bool = False) -> None:
    config.ensure_dirs()
    out_dir = config.SOURCES_DIR
    manifest_path = out_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    records: list[dict] = []
    seen_hashes: dict[str, str] = {}

    root = config.CAREERS_FOLDER
    if not root.exists():
        sys.exit(f"Careers folder not found: {root}")

    files = sorted(p for p in root.rglob("*") if p.is_file())
    for path in files:
        rel = path.relative_to(root)
        if any(part in SKIP_DIRS for part in rel.parts) or path.suffix.lower() in SKIP_SUFFIXES:
            continue
        extractor = EXTRACTORS.get(path.suffix.lower())
        if not extractor:
            print(f"  - skip (no extractor): {rel}")
            continue
        try:
            text, meta = extractor(path)
        except Exception as exc:  # noqa: BLE001
            print(f"  ! failed: {rel}: {exc}")
            continue

        kind, voice, note = kind_for(rel)
        stem = path.stem
        sid = f"{kind}--{slugify(stem)}--{path.suffix.lower().lstrip('.')}"
        # The same file name in two folders (e.g. a root-level copy of a cover letter) must not
        # share an id; disambiguate with the folder slug.
        if any(r["meta"]["id"] == sid for r in records):
            sid = f"{sid}--{slugify(rel.parent.as_posix()) or 'root'}"
        text_hash = sha256_text(text)
        duplicate_of = seen_hashes.get(text_hash)
        if not duplicate_of:
            seen_hashes[text_hash] = sid

        record = {
            "id": sid,
            "kind": kind,
            "title": stem,
            "employer": guess_employer(stem),
            "source_path": str(path),
            "relative_path": rel.as_posix(),
            "format": path.suffix.lower().lstrip("."),
            "file_modified": dt.datetime.fromtimestamp(path.stat().st_mtime).date().isoformat(),
            "sha256_file": sha256_file(path),
            "sha256_text": text_hash,
            "words": len(text.split()),
            "voice_corpus": voice and not duplicate_of,
            "duplicate_of": duplicate_of,
            "notes": note,
            **meta,
        }
        records.append({"meta": record, "text": text})

    # Format pairs: the same document as .docx and .pdf with slightly different extraction.
    # Prefer the .docx (cleaner text) and mark the .pdf as a format duplicate.
    by_key: dict[tuple[str, str], list[dict]] = {}
    for r in records:
        m = r["meta"]
        by_key.setdefault((m["kind"], slugify(re.sub(r"\s*\(\d\)$", "", m["title"]))), []).append(r)
    for group in by_key.values():
        if len(group) < 2:
            continue
        preferred = sorted(group, key=lambda r: (r["meta"]["format"] != "docx", r["meta"]["relative_path"]))[0]
        for r in group:
            if r is not preferred and not r["meta"]["duplicate_of"]:
                r["meta"]["duplicate_of"] = preferred["meta"]["id"]
                r["meta"]["voice_corpus"] = False
                r["meta"]["notes"] = (r["meta"]["notes"] or "") + " format duplicate; the .docx version is canonical"

    wl = ingest_work_log(config.WORK_LOG_XLSX)
    if wl:
        records.append(wl)

    # Write files
    written = skipped = 0
    for r in records:
        m = r["meta"]
        m = {k: v for k, v in m.items() if v not in (None, "", [])}
        m["extracted_at"] = dt.datetime.now().isoformat(timespec="seconds")
        target = out_dir / f"{m['id']}.md"
        key = m["id"]
        prev = manifest.get(key)
        if prev and not force and prev.get("sha256_text") == m.get("sha256_text", m.get("sha256")) and target.exists():
            skipped += 1
            continue
        target.write_text(front_matter(m) + "\n" + r["text"].strip() + "\n", encoding="utf-8")
        manifest[key] = {"path": str(target.relative_to(config.REPO_ROOT)), "sha256_text": m.get("sha256_text", m.get("sha256")), "extracted_at": m["extracted_at"]}
        written += 1

    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    # Hand-written sources (e.g. the transcribed transcript, the live LinkedIn read) are not generated
    # by this script but belong in the index; read their front matter.
    generated = {Path(v["path"]).name for v in manifest.values()}
    for extra in sorted(out_dir.glob("*.md")):
        if extra.name in generated or extra.name == "INDEX.md":
            continue
        raw = extra.read_text(encoding="utf-8")
        if not raw.startswith("---"):
            continue
        try:
            fm = yaml.safe_load(raw.split("---", 2)[1]) or {}
        except yaml.YAMLError:
            continue
        fm.setdefault("id", extra.stem)
        fm.setdefault("kind", "manual")
        fm["words"] = len(raw.split())
        fm["notes"] = (fm.get("notes") or "") + " (hand-written source)"
        records.append({"meta": fm, "text": ""})

    # INDEX.md
    lines = [
        "# Sources index",
        "",
        f"Generated {dt.date.today().isoformat()} by `python -m pipeline.ingest`. {len(records)} sources; "
        f"{sum(1 for r in records if r['meta'].get('duplicate_of'))} marked as duplicates; "
        f"{sum(1 for r in records if r['meta'].get('voice_corpus'))} in the voice corpus.",
        "",
        "| id | kind | employer | date | words | voice | duplicate of | notes |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for r in sorted(records, key=lambda r: (r["meta"]["kind"], r["meta"]["id"])):
        m = r["meta"]
        date = m.get("doc_modified") or m.get("file_modified") or m.get("date_range") or ""
        lines.append(
            f"| `{m['id']}` | {m['kind']} | {m.get('employer') or ''} | {date} | {m.get('words', '')} | "
            f"{'yes' if m.get('voice_corpus') else ''} | {m.get('duplicate_of') or ''} | {(m.get('notes') or '').strip()} |"
        )
    (out_dir / "INDEX.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {written}, skipped {skipped} unchanged, total {len(records)} sources -> {out_dir}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--force", action="store_true", help="re-extract everything")
    args = ap.parse_args()
    ingest(force=args.force)
